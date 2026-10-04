"""Build the Stage 2 + Stage 3 viva deck (16:9) with python-pptx."""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "results" / "figures"
OUT = ROOT.parent / "DR_Grading_Stage2_3_Presentation.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
CORAL = RGBColor(0xE7, 0x6F, 0x51)
INK = RGBColor(0x22, 0x2B, 0x36)
GREY = RGBColor(0x6B, 0x75, 0x80)
LIGHT = RGBColor(0xF1, 0xF4, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333
state = {"n": 0}


def text(slide, x, y, w, h, content, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font="Calibri"):
    """content: str or list of str / (str, dict) runs-per-paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        opts = {}
        if isinstance(p, tuple):
            p, opts = p
        r = para.add_run()
        r.text = p
        r.font.name = font
        r.font.size = Pt(opts.get("size", size))
        r.font.bold = opts.get("bold", bold)
        r.font.color.rgb = opts.get("color", color)
        para.space_after = Pt(opts.get("space", 6))
    return tb


def bullets(slide, x, y, w, h, items, size=18, color=INK):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        level = 0
        if isinstance(item, tuple):
            item, level = item
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = para.add_run()
        r.text = ("•  " if level == 0 else "–  ") + item
        r.font.name = "Calibri"
        r.font.size = Pt(size - 2 * level)
        r.font.color.rgb = color if level == 0 else GREY
        para.level = level
        para.space_after = Pt(8 if level == 0 else 4)
        if level:
            para.left_indent = Inches(0.35)
    return tb


def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def card(slide, x, y, w, h, title, body, accent=TEAL, title_size=16, body_size=14):
    box(slide, x, y, w, h, LIGHT)
    box(slide, x, y, 0.09, h, accent, shape=MSO_SHAPE.RECTANGLE)
    text(slide, x + 0.22, y + 0.12, w - 0.35, 0.45, title, size=title_size, bold=True, color=accent)
    text(slide, x + 0.22, y + 0.55, w - 0.35, h - 0.6, body, size=body_size, color=INK)


def stat(slide, x, y, w, value, label, color=NAVY):
    text(slide, x, y, w, 0.8, value, size=40, bold=True, color=color, align=PP_ALIGN.CENTER)
    text(slide, x, y + 0.8, w, 0.6, label, size=13, color=GREY, align=PP_ALIGN.CENTER)


def image(slide, path, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def table(slide, x, y, w, rows, col_w, size=13, header_fill=NAVY, row_h=0.36, highlight=None):
    nr, nc = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = shp.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            r = tf.paragraphs[0].add_run()
            r.text = str(val)
            r.font.name = "Calibri"
            r.font.size = Pt(size)
            tf.paragraphs[0].alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
                r.font.color.rgb = WHITE
                r.font.bold = True
            else:
                hl = highlight and i in highlight
                cell.fill.fore_color.rgb = highlight[i] if hl else (LIGHT if i % 2 else WHITE)
                r.font.color.rgb = INK
                r.font.bold = bool(hl)
    return tbl


def arrow(slide, x1, y1, x2, y2, color=GREY):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2.25)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def new_slide(title, kicker=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    state["n"] += 1
    box(s, 0, 0, W, 1.05, NAVY, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        text(s, 0.55, 0.1, 9, 0.35, kicker.upper(), size=12, bold=True, color=RGBColor(0x9F, 0xD8, 0xCF))
    text(s, 0.55, 0.36 if kicker else 0.2, 12.2, 0.65, title, size=28, bold=True, color=WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.55, 7.08, 9, 0.3, "Diabetic Retinopathy Grading · Reproduction of Chilukoti et al. (2024) · Stage 2 & 3",
         size=10, color=GREY)
    text(s, W - 1.2, 7.08, 0.7, 0.3, str(state["n"]), size=10, color=GREY, align=PP_ALIGN.RIGHT)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def blank():
    state["n"] += 1
    return prs.slides.add_slide(BLANK)


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


PDF = "--pdf" in sys.argv
if PDF:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import pdf_backend as pb

    deck = pb.Deck(OUT.with_suffix(".pdf"))
    _ALIGN = {PP_ALIGN.LEFT: "left", PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}
    _hex = pb.hexc

    def text(slide, x, y, w, h, content, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, font="Calibri"):
        paras = []
        for p in (content if isinstance(content, list) else [content]):
            opts = {}
            if isinstance(p, tuple):
                p, opts = p
            paras.append((p, opts.get("size", size), _hex(opts.get("color", color)), opts.get("bold", bold),
                          opts.get("space", 6), 0))
        pb.draw_paragraphs(slide, x, y, w, h, paras, _ALIGN.get(align, "left"),
                           "middle" if anchor == MSO_ANCHOR.MIDDLE else "top")

    def bullets(slide, x, y, w, h, items, size=18, color=INK):
        paras = []
        for item in items:
            level = 0
            if isinstance(item, tuple):
                item, level = item
            paras.append((("•  " if level == 0 else "–  ") + item, size - 2 * level,
                          _hex(color if level == 0 else GREY), False, 8 if level == 0 else 4, 0.35 * level))
        pb.draw_paragraphs(slide, x, y, w, h, paras)

    def box(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
        pb.rect(slide, x, y, w, h, _hex(fill), rounded=shape == MSO_SHAPE.ROUNDED_RECTANGLE,
                edge=_hex(line) if line else None)

    def image(slide, path, x, y, w=None, h=None):
        pb.picture(slide, path, x, y, w, h)

    def table(slide, x, y, w, rows, col_w, size=13, header_fill=NAVY, row_h=0.36, highlight=None):
        hl = {k: _hex(v) for k, v in (highlight or {}).items()}
        pb.grid(slide, x, y, rows, col_w, size, _hex(header_fill), row_h, hl, (_hex(INK), _hex(LIGHT), _hex(WHITE)))

    def arrow(slide, x1, y1, x2, y2, color=GREY):
        pb.line_arrow(slide, x1, y1, x2, y2, _hex(color))

    def blank():
        state["n"] += 1
        return deck.new()

    def notes(slide, txt):
        pass

    def new_slide(title, kicker=None, notes=""):
        s = blank()
        box(s, 0, 0, W, 1.05, NAVY, shape=MSO_SHAPE.RECTANGLE)
        if kicker:
            text(s, 0.55, 0.1, 9, 0.35, kicker.upper(), size=12, bold=True, color=RGBColor(0x9F, 0xD8, 0xCF))
        text(s, 0.55, 0.36 if kicker else 0.2, 12.2, 0.65, title, size=28, bold=True, color=WHITE,
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, 0.55, 7.08, 9, 0.3, "Diabetic Retinopathy Grading · Reproduction of Chilukoti et al. (2024) · "
             "Stage 2 & 3", size=10, color=GREY)
        text(s, W - 1.2, 7.08, 0.7, 0.3, str(state["n"]), size=10, color=GREY, align=PP_ALIGN.RIGHT)
        return s

# 1. Title -------------------------------------------------------------------------------------------
s = blank()
box(s, 0, 0, W, 7.5, NAVY, shape=MSO_SHAPE.RECTANGLE)
box(s, 0.7, 2.05, 0.12, 2.3, TEAL, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.05, 1.0, 11, 0.5, "CV RESEARCH-REPRODUCTION PROJECT  ·  STAGE 2 & STAGE 3", size=14, bold=True,
     color=RGBColor(0x9F, 0xD8, 0xCF))
text(s, 1.05, 1.95, 12.0, 1.8, "Reproducing Transfer-Learning & Snapshot-Ensemble\nDiabetic Retinopathy Grading on APTOS 2019",
     size=33, bold=True, color=WHITE)
text(s, 1.05, 3.75, 11, 0.6, "Base paper: Chilukoti et al., BMC Medical Informatics and Decision Making 24:37 (2024)",
     size=18, color=RGBColor(0xD6, 0xE2, 0xF0))
text(s, 1.05, 5.0, 11, 0.5, "Shivansh  ·  Divyanshi  ·  Vivek  ·  Shubham", size=22, bold=True, color=WHITE)
text(s, 1.05, 5.6, 11, 0.5, "github.com/sshivanshg/dr-grading-reproduction", size=14,
     color=RGBColor(0x9F, 0xD8, 0xCF))
notes(s, (
    "Introduce the team and the task: five-class DR grading on APTOS 2019. Today we present Stage 2 (choosing and "
    "auditing a base paper) and Stage 3 (reproducing its baseline and explaining the gap)."))

# 2. Where we are --------------------------------------------------------------------------------------
s = new_slide("Where we are in the project", "Roadmap",
              "Stage 1 was the literature survey. Today covers Stage 2 and Stage 3. Stage 4 (hypothesis) is next.")
steps = [("1", "Literature survey", "10+ papers, gaps identified", GREY),
         ("2", "Pick base paper", "Reproducibility checklist", TEAL),
         ("3", "Reproduce", "Re-run method, explain gap", TEAL),
         ("4", "Hypothesis", "One change + prediction", GREY),
         ("5", "Implement & ablate", "Measure what moved", GREY),
         ("6", "Analyse & write", "IEEE paper + defence", GREY)]
for i, (num, t, sub, col) in enumerate(steps):
    x = 0.55 + i * 2.08
    box(s, x, 1.7, 1.9, 2.1, col if col == TEAL else LIGHT)
    fg = WHITE if col == TEAL else INK
    text(s, x, 1.85, 1.9, 0.7, num, size=34, bold=True, color=fg, align=PP_ALIGN.CENTER)
    text(s, x + 0.05, 2.55, 1.8, 0.6, t, size=14, bold=True, color=fg, align=PP_ALIGN.CENTER)
    text(s, x + 0.05, 3.1, 1.8, 0.65, sub, size=11, color=fg if col == TEAL else GREY, align=PP_ALIGN.CENTER)
text(s, 2.63, 3.95, 4.0, 0.4, "▲ TODAY", size=14, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
card(s, 0.55, 4.6, 6.0, 2.2, "Stage 1 take-aways (our literature review)",
     ["Class imbalance remains unresolved: Severe NPDR & PDR are rare",
      "Accuracy can be misleading; per-class recall and QWK are needed",
      "Ordinal nature of grades (0→4) is under-used"], accent=NAVY, body_size=14)
card(s, 6.8, 4.6, 6.0, 2.2, "Research question we are building towards",
     ["Can class-balanced, ordinal-aware learning improve Severe NPDR and PDR recognition while reducing large "
      "grading errors?", ("Stages 2–3 build the reproducible baseline to test it on.", {"color": GREY, "size": 13})],
     accent=TEAL, body_size=14)

# 3. Task & dataset ------------------------------------------------------------------------------------
s = new_slide("Task: five-class DR grading on APTOS 2019", "Problem",
              "APTOS 2019: 3,662 fundus photos from Aravind Eye Hospital, graded 0-4. Heavily imbalanced: No DR is "
              "half the data; Severe is only 5%. Grades are ordinal, so QWK is the standard metric.")
image(s, FIG / "class_distribution.png", 0.45, 1.3, w=6.9)
card(s, 7.65, 1.35, 5.15, 1.55, "Dataset", ["3,662 colour fundus photographs, Aravind Eye Hospital (rural India)",
                                            "Graded 0–4 by clinicians (ICDR scale)"], accent=NAVY, body_size=14)
card(s, 7.65, 3.05, 5.15, 1.55, "Split (stratified, seed 42)", ["70 / 15 / 15  →  2,562 train · 550 val · 550 test",
                                                               "Same proportions the paper used for EyePACS"],
     accent=NAVY, body_size=14)
card(s, 7.65, 4.75, 5.15, 2.05, "Why it is hard", ["Severe NPDR = 5%, PDR = 8% of images",
                                                   "Adjacent grades look alike (Mild↔Moderate, Severe↔PDR)",
                                                   "Grades are ordinal, so error size matters → QWK"],
     accent=CORAL, body_size=14)

# 4. Base paper --------------------------------------------------------------------------------------------
s = new_slide("Base paper: Chilukoti et al. (2024)", "Stage 2 · Selection",
              "Walk through the pipeline left to right. The two claimed contributions are CLAHE + blur preprocessing and "
              "the computationally cheap snapshot ensemble. Reported QWK: 0.954 single model, 0.967 with EyePACS "
              "pre-training plus the ensemble.")
stages = [("Fundus image", "APTOS 2019"), ("CLAHE + Gaussian blur", "contrast + denoise, 300 px"),
          ("EfficientNet-B3", "ImageNet-pretrained, fine-tuned"), ("3-layer head", "512 → 512 → 5, dropout 0.5/0.25"),
          ("Snapshot vote", "10 saved epochs, majority")]
for i, (t, sub) in enumerate(stages):
    x = 0.55 + i * 2.5
    box(s, x, 1.55, 2.1, 1.45, NAVY if i in (1, 4) else LIGHT)
    fg = WHITE if i in (1, 4) else INK
    text(s, x + 0.08, 1.68, 1.94, 0.6, t, size=15, bold=True, color=fg, align=PP_ALIGN.CENTER)
    text(s, x + 0.08, 2.3, 1.94, 0.6, sub, size=11, color=fg if i in (1, 4) else GREY, align=PP_ALIGN.CENTER)
    if i < 4:
        arrow(s, x + 2.12, 2.27, x + 2.48, 2.27)
text(s, 0.55, 3.08, 12, 0.35, "Dark boxes = the paper's two claimed contributions", size=12, color=GREY)
table(s, 0.55, 3.7, 6.3, [["Reported on APTOS 2019", "QWK"],
                          ["ImageNet init, single model", "0.954"],
                          ["EyePACS init, single model", "0.956"],
                          ["EyePACS init + snapshot ensemble", "0.967"]],
      [4.8, 1.5], size=15, row_h=0.45, highlight={1: RGBColor(0xD8, 0xEF, 0xEC)})
text(s, 0.55, 5.65, 6.3, 0.9, "Our target: 0.954 (ImageNet init). EyePACS pre-training needs 35k extra images, "
     "beyond our compute.", size=14, color=GREY)
card(s, 7.2, 3.7, 5.6, 3.05, "Training recipe stated in the paper",
     ["Adam, learning rate 1e-3, weight decay 1e-4", "Gradient clipping 0.1",
      "Strategy 2: CLAHE + blur at 300 × 300, weights saved at 10 intervals",
      "QWK as the primary metric"], accent=TEAL, body_size=14)

# 5. Why this paper ----------------------------------------------------------------------------------------
s = new_slide("Why we chose this paper", "Stage 2 · Selection",
              "Selection logic: it is on our chosen dataset, uses an ordinal metric, fits free hardware, and leaves "
              "exactly the gaps from our literature review open, which makes it a clean base for a one-variable hypothesis.")
reasons = [("Same dataset & task", "Five-class grading on APTOS 2019, the course-recommended dataset", NAVY),
           ("Ordinal metric", "Uses QWK as primary metric, matching our Stage 1 finding", NAVY),
           ("Feasible compute", "Transfer learning from one pretrained CNN, no heavy ensemble of models", NAVY),
           ("Leaves our gaps open", "No imbalance-aware or ordinal loss, so there's room for one clean change", TEAL),
           ("Strong claim to test", "0.954–0.967 QWK is state-of-the-art level, worth verifying", TEAL),
           ("Clear method text", "Architecture and optimiser are described well enough to re-implement", TEAL)]
for i, (t, b, col) in enumerate(reasons):
    r, c = divmod(i, 3)
    card(s, 0.55 + c * 4.15, 1.45 + r * 2.75, 3.95, 2.5, t, [b], accent=col, title_size=18, body_size=15)

# 6. Reproducibility checklist -----------------------------------------------------------------------------
s = new_slide("Reproducibility checklist", "Stage 2 · Checklist",
              "Be upfront: no public code exists, so everything is our own re-implementation. Several key details "
              "are missing; slide 7 lists our choices for each.")
rows = [["Checklist item", "Our answer", "Status"],
        ["Public code?", "No repository linked or found; re-implemented from the paper text", "✗"],
        ["Dataset public & size?", "APTOS 2019, 3,662 images, ~8.9 GB (full-res mirror on Hugging Face)", "✓"],
        ["Fits free GPU?", "Expected yes on Colab T4 at 300 px; locally (M2, 8 GB) only at 224 px", "~"],
        ["Code usable / last commit?", "N/A; our code: PyTorch 2.14 + timm 1.0.30, resumable training", "—"],
        ["Metric defined?", "QWK (primary) + accuracy, precision, recall, F1, AUC", "✓"],
        ["Split defined?", "Only for EyePACS (70/15/15); not for APTOS", "~"],
        ["Hyper-parameters?", "Optimiser, lr, decay, clipping given; batch size, loss, epochs, CLAHE params not", "~"]]
table(s, 0.55, 1.4, 12.25, rows, [2.9, 8.35, 1.0], size=14, row_h=0.6,
      highlight={1: RGBColor(0xFB, 0xE3, 0xDC)})
text(s, 0.55, 6.35, 12.2, 0.5, "✓ available    ~ partially specified    ✗ missing", size=13, color=GREY)

# 7. Under-specified details -------------------------------------------------------------------------------
s = new_slide("Filling the gaps: our implementation choices", "Stage 2 · Checklist",
              "Each missing detail gets a documented, defensible choice. Also mention the red flags: some reported "
              "numbers are internally inconsistent, so 0.954 may itself be unreliable.")
rows = [["Not stated in paper", "Our choice", "Why"],
        ["Loss function", "Cross-entropy", "Softmax output implies it"],
        ["“Gradient clipping 0.1”", "Global L2-norm clip at 0.1", "Most common reading"],
        ["CLAHE settings", "Clip 2.0, 8×8 tiles, on LAB luminance", "Standard; keeps colour"],
        ["Blur", "5 × 5 Gaussian, after CLAHE", "Order as written in paper"],
        ["APTOS split", "Stratified 70/15/15, seed 42", "Paper's EyePACS proportions"],
        ["Final or best model?", "Report both + both ensembles", "Avoid cherry-picking"]]
table(s, 0.55, 1.4, 7.6, rows, [2.35, 3.0, 2.25], size=13, row_h=0.55)
card(s, 8.45, 1.4, 4.35, 5.35, "Red flags in the paper",
     ["Table 3: precision = recall = QWK = 0.867 exactly",
      "Constant predictors report macro AUC 0.196 (a constant predictor gives 0.5)",
      "Table 6: accuracy, precision, recall, F1 all within 0.01 on imbalanced data",
      "APTOS split unspecified → possible leakage",
      ("We treat 0.954 as a claim to test, not ground truth.", {"bold": True, "color": CORAL, "size": 14})],
     accent=CORAL, body_size=14)

# 8. Data audit --------------------------------------------------------------------------------------------
s = new_slide("Data audit: APTOS contains duplicate images", "Stage 3 · Protocol",
              "We hashed every image's raw bytes. 123 groups are exact copies; in 30 of them the same picture has two "
              "different grades. Under a random split like the paper's, 29 test images have a copy in training. "
              "We later show this has almost no effect on QWK.")
stat(s, 0.55, 1.55, 3.0, "123", "groups of byte-identical images")
stat(s, 3.65, 1.55, 3.0, "30", "groups with conflicting grades", color=CORAL)
stat(s, 6.75, 1.55, 3.0, "29 / 550", "test images with a copy in train")
stat(s, 9.85, 1.55, 3.0, "±0.001", "QWK change when removed", color=TEAL)
card(s, 0.55, 3.45, 6.0, 3.3, "What we did",
     ["SHA-1 hash of every raw image file",
      "Perceptual hashing rejected: every fundus is a bright disc on black, giving thousands of false matches",
      "Re-scored each model on the 521 test images without a training copy"], accent=NAVY, body_size=15)
card(s, 6.8, 3.45, 6.0, 3.3, "What it means",
     ["Leakage does not explain our gap (QWK 0.827 → 0.828)",
      "Leaked images were predicted worse (65.5% acc.), because of conflicting labels",
      "Label noise exists in APTOS itself (e.g. the same image graded Moderate and PDR)"], accent=TEAL, body_size=15)

# 9. Hardware & deviations ---------------------------------------------------------------------------------
s = new_slide("Implementation under hardware limits", "Stage 3 · Protocol",
              "We trained locally on an Apple M2 with 8 GB of shared memory. B3 at 300 px swapped to disk even at batch 8 "
              "(25-45 s per step), so we used 224 px. Everything else follows the paper. The exact protocol is one "
              "flag change and ships as a Colab notebook.")
rows = [["Setting", "Paper", "Ours", "Reason"],
        ["Backbone / head / optimiser", "EffNet-B3, 3-layer, Adam 1e-3", "Identical", "—"],
        ["Pre-processing", "CLAHE + blur", "Identical", "—"],
        ["Resolution", "300 × 300", "224 × 224", "300 px exceeds 8 GB memory"],
        ["Epochs / snapshots", "60 / 10", "30 / 10 (every 3)", "~4 min per epoch locally"],
        ["Batch size", "not stated", "8", "largest without swapping"],
        ["EyePACS pre-training", "yes (for 0.967)", "no", "35k images, out of budget"]]
table(s, 0.55, 1.4, 8.3, rows, [2.45, 2.15, 1.75, 1.95], size=13, row_h=0.55,
      highlight={3: RGBColor(0xFB, 0xE3, 0xDC), 4: RGBColor(0xFB, 0xE3, 0xDC)})
card(s, 9.15, 1.4, 3.65, 2.45, "Benchmarked on M2", ["B3 @ 300 px: swaps, 25–45 s / step",
                                                       "B3 @ 224 px, bs 8: 2.8 GB, 12.4 img/s",
                                                       "Mixed precision: no speed-up"], accent=NAVY, body_size=13)
card(s, 9.15, 4.05, 3.65, 2.7, "Engineering", ["Pre-processed cache (memory-mapped)",
                                                "Resumable checkpoints every epoch",
                                                "Seeded: bit-identical across launches",
                                                "Colab notebook for exact protocol"], accent=TEAL, body_size=13)

# 10. Main results -----------------------------------------------------------------------------------------
s = new_slide("Results: reproduction vs. paper", "Stage 3 · Results",
              "Headline: with the paper's settings we get 0.827, a gap of about 0.13. Changing only the learning rate "
              "gives 0.887. Point at the Severe column: zero with the paper's settings.")
rows = [["Configuration", "Prediction", "QWK", "Acc.", "Macro-F1", "AUC", "Severe rec.", "PDR rec."],
        ["Paper settings (lr 1e-3)", "single model", "0.827", "0.771", "0.473", "0.897", "0.00", "0.30"],
        ["Paper settings (lr 1e-3)", "10-snapshot vote", "0.788", "0.767", "0.412", "0.913", "0.00", "0.00"],
        ["Only change: lr 1e-4", "single model", "0.887", "0.815", "0.630", "0.894", "0.31", "0.55"],
        ["Only change: lr 1e-4", "10-snapshot vote", "0.887", "0.822", "0.613", "0.930", "0.10", "0.59"],
        ["Paper (reported)", "single / ensemble", "0.954 / 0.967", "0.942", "0.944", "—", "—", "—"]]
table(s, 0.55, 1.4, 12.25, rows, [2.6, 1.95, 1.65, 1.0, 1.2, 1.0, 1.4, 1.45], size=14, row_h=0.55,
      highlight={3: RGBColor(0xD8, 0xEF, 0xEC), 4: RGBColor(0xD8, 0xEF, 0xEC)})
stat(s, 0.55, 5.0, 3.0, "−0.127", "QWK gap with paper settings", color=CORAL)
stat(s, 3.65, 5.0, 3.0, "+0.060", "QWK from lr change alone", color=TEAL)
stat(s, 6.75, 5.0, 3.0, "0 → 0.31", "Severe NPDR recall", color=TEAL)
stat(s, 9.85, 5.0, 3.0, "48 → 24", "errors of ≥ 2 grades", color=TEAL)

# 11. Training dynamics ------------------------------------------------------------------------------------
s = new_slide("Why: training at lr 1e-3 is unstable", "Stage 3 · Analysis",
              "Left: paper settings. Validation QWK crashes at epochs 2, 5, 10, 17 and 23; at epoch 5 the model "
              "predicts almost one grade. Loss is still 0.61 after 30 epochs. Right: lr 1e-4, smooth, loss ~0.3.")
image(s, FIG / "s2_faithful_curves.png", 0.4, 1.35, w=6.2)
image(s, FIG / "s2_lr1e4_curves.png", 6.75, 1.35, w=6.2)
text(s, 0.4, 3.65, 6.2, 0.4, "Paper settings: lr 1e-3", size=16, bold=True, color=CORAL, align=PP_ALIGN.CENTER)
text(s, 6.75, 3.65, 6.2, 0.4, "Only change: lr 1e-4", size=16, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
card(s, 0.55, 4.35, 6.0, 2.4, "Symptoms",
     ["Validation QWK collapses 5 times (down to −0.006 at epoch 5)",
      "Training loss still 0.61 after 30 epochs: under-fitting, not over-fitting"], accent=CORAL, body_size=15)
card(s, 6.8, 4.35, 6.0, 2.4, "Explanation",
     ["Adam at 1e-3 over-writes pretrained features in full fine-tuning",
      "Small batch (8) adds noisy batch-norm statistics",
      "Paper's own Table 3 shows VGG/ResNet collapsing to class 0"], accent=TEAL, body_size=15)

# 12. Confusion matrices -----------------------------------------------------------------------------------
s = new_slide("QWK hides minority-grade failure", "Stage 3 · Analysis",
              "Left: with the paper's settings Severe is never predicted; 15 of 29 go to Moderate and 13 to PDR. QWK "
              "stays at 0.83 because these errors are adjacent. This is exactly the 'accuracy can be misleading' gap "
              "from our literature review.")
image(s, FIG / "s2_faithful_best_val_cm.png", 0.45, 1.25, h=4.1)
image(s, FIG / "s2_lr1e4_best_val_cm.png", 5.35, 1.25, h=4.1)
text(s, 0.45, 5.4, 4.8, 0.4, "Paper settings (lr 1e-3)", size=15, bold=True, color=CORAL, align=PP_ALIGN.CENTER)
text(s, 5.35, 5.4, 4.8, 0.4, "lr 1e-4", size=15, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
card(s, 10.35, 1.35, 2.45, 5.4, "Read-out",
     ["Severe NPDR: 0 / 29 correct at lr 1e-3", "27 / 44 PDR called Moderate",
      "Mild mostly → Moderate", "Errors are adjacent, so QWK stays ≈ 0.83",
      ("Per-class recall is essential", {"bold": True, "color": CORAL, "size": 13})], accent=CORAL, body_size=13)
text(s, 0.45, 6.0, 9.7, 0.8, "At lr 1e-4: Severe 9 / 29 correct, PDR 24 / 44; large errors halve. Severe remains the "
     "hardest grade, which is the target for Stage 4.", size=14, color=GREY)

# 13. Snapshot ensemble ------------------------------------------------------------------------------------
s = new_slide("The snapshot ensemble only works when training is stable", "Stage 3 · Analysis",
              "Each dot is one saved epoch. Paper settings (coral): only the last snapshot is good, so the vote of "
              "ten is dragged down to 0.788. lr 1e-4 (teal): every snapshot is 0.83-0.88 and the vote, 0.887, beats "
              "every individual snapshot. Snapshot ensembles assume comparably good members.")
image(s, FIG / "snapshot_qwk.png", 0.4, 1.3, w=7.6)
card(s, 8.3, 1.4, 4.5, 2.55, "Paper settings (lr 1e-3)",
     ["Snapshot QWK 0.65 – 0.79, last one 0.83", "Majority vote = 0.788, worse than single model",
      "Severe & PDR recall in vote: 0.00"], accent=CORAL, body_size=14)
card(s, 8.3, 4.15, 4.5, 2.6, "Stable training (lr 1e-4)",
     ["Every snapshot 0.83 – 0.88", "Vote = 0.887, above all snapshots",
      "Ensemble helps only with good, diverse members (Huang et al., 2017)"], accent=TEAL, body_size=14)

# 14. Gap analysis -----------------------------------------------------------------------------------------
s = new_slide("Explaining the reproduction gap (0.827 vs 0.954)", "Stage 3 · Gap analysis",
              "Order by strength of evidence: the learning rate is proven experimentally; resolution and missing "
              "details are plausible; leakage is ruled out; the paper's own numbers are questionable.")
rows = [["Candidate cause", "Evidence", "Verdict"],
        ["Learning rate 1e-3 (paper's setting)", "Changing only lr → +0.060 QWK, stable training", "Main cause"],
        ["224 px / 30 epochs (hardware)", "Small lesions less visible; not tested yet (Colab notebook ready)", "Likely part"],
        ["Unstated details (split, batch, loss)", "Each can move QWK a few points on 550 test images", "Possible"],
        ["Duplicate leakage", "QWK 0.827 → 0.828 without the 29 leaked images", "Ruled out"],
        ["Reliability of reported numbers", "Internally inconsistent tables (slide 7)", "Open question"]]
tbl = table(s, 0.55, 1.45, 12.25, rows, [3.9, 6.25, 2.1], size=15, row_h=0.72,
            highlight={1: RGBColor(0xD8, 0xEF, 0xEC), 4: RGBColor(0xF1, 0xF4, 0xF8)})
text(s, 0.55, 6.0, 12.2, 0.8, "The single change of learning rate closes about half the gap; the rest is consistent "
     "with our lower resolution and shorter schedule.", size=16, bold=True, color=NAVY)

# 15. Failure cases ----------------------------------------------------------------------------------------
s = new_slide("Failure cases", "Stage 3 · Analysis",
              "The ten most confidently wrong predictions of the paper-settings model. Mostly PDR predicted as Mild or "
              "Moderate, including one with a large haemorrhage. Several are dim, low-contrast or cropped.")
image(s, FIG / "s2_faithful_failures.png", 1.28, 1.3, h=4.9)
text(s, 0.55, 6.35, 12.2, 0.6, "Mostly PDR → Mild/Moderate (incl. a large pre-retinal haemorrhage); several are dim, "
     "low-contrast or cropped. Neovascularisation is hard to see at 224 px.", size=15, color=INK)

# 16. Limitations ------------------------------------------------------------------------------------------
s = new_slide("Limitations (and what we are honest about)", "Stage 3 · Limitations",
              "Raise these before the examiners do.")
card(s, 0.55, 1.4, 6.0, 5.35, "Experimental",
     ["Single dataset (APTOS), single seed, no external validation",
      "224 px / 30 epochs instead of 300 px / 60 epochs",
      "EyePACS-pretrained setting (0.967) not reproduced",
      "Only 29 Severe and 44 PDR test images: 1 image ≈ 3 recall points",
      "Paper's strategy 1 (150 px) skipped: reported on EyePACS only"], accent=CORAL, body_size=15)
card(s, 6.8, 1.4, 6.0, 5.35, "Process",
     ["Headline run interrupted (disk full) and resumed from epoch 13 checkpoint",
      "The lr 1e-4 run is our diagnostic, not the paper's protocol",
      "No public code, so some choices are our interpretation",
      ("No clinical claims: this is a reproduction study", {"bold": True, "color": NAVY, "size": 15})],
     accent=NAVY, body_size=15)

# 17. Next: Stage 4 ----------------------------------------------------------------------------------------
s = new_slide("Next: Stage 4 hypothesis (predicted before coding)", "Looking ahead",
              "One variable (the loss), a direction, a magnitude, and a mechanism, as the course requires. Built on the "
              "stable lr 1e-4 baseline.")
card(s, 0.55, 1.4, 12.25, 1.65, "Hypothesis",
     ["Replacing cross-entropy with a class-balanced, ordinal-aware loss on the stable lr 1e-4 baseline will improve "
      "Severe NPDR and PDR recognition and reduce large grading errors."], accent=NAVY, title_size=18, body_size=17)
card(s, 0.55, 3.25, 6.0, 3.5, "Mechanism",
     ["Class weights raise the gradient contribution of rare grades",
      "Ordinal term penalises errors by distance from the true grade",
      "Targets both gaps from our literature review"], accent=TEAL, body_size=15)
rows = [["Metric", "Baseline", "Predicted"],
        ["Severe recall", "0.31", "+0.10 to +0.20"],
        ["Minority recall", "0.43", "increase"],
        ["Errors ≥ 2 grades", "24", "< 24"],
        ["QWK", "0.887", "0 to +0.02"],
        ["Accuracy", "0.815", "may dip 1–2 pts"]]
table(s, 6.8, 3.25, 6.0, rows, [2.4, 1.5, 2.1], size=15, row_h=0.58)

# 18. Summary ----------------------------------------------------------------------------------------------
s = new_slide("Summary", "Take-aways",
              "Close with the four messages and invite questions. Point to the repo for code, logs and figures.")
msgs = [("1", "Reproduced from scratch", "No public code. With the paper's settings: QWK 0.827 vs 0.954 claimed.", NAVY),
        ("2", "Found the main cause", "lr 1e-3 makes training unstable; lr 1e-4 alone gives 0.887.", TEAL),
        ("3", "Ensemble claim is conditional", "Snapshot vote hurts when unstable (0.788) and helps when stable.", TEAL),
        ("4", "QWK is not enough", "Severe NPDR was never detected at QWK 0.83: report per-class recall.", CORAL)]
for i, (num, t, b, col) in enumerate(msgs):
    y = 1.4 + i * 1.3
    box(s, 0.55, y, 0.95, 1.1, col)
    text(s, 0.55, y, 0.95, 1.1, num, size=30, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.75, y + 0.05, 10.8, 0.5, t, size=20, bold=True, color=col)
    text(s, 1.75, y + 0.52, 10.8, 0.5, b, size=16, color=INK)
text(s, 0.55, 6.6, 12.2, 0.4, "Code, logs, figures & report: github.com/sshivanshg/dr-grading-reproduction   ·   "
     "Questions?", size=15, bold=True, color=NAVY)

if PDF:
    deck.close()
    print("saved", OUT.with_suffix(".pdf"), "slides:", state["n"])
else:
    prs.save(OUT)
    print("saved", OUT, "slides:", len(prs.slides))
